"""DOCX parser (python-docx): paragraph styles drive the structure."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import docx
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph

from .base import BlockKind, ParsedDocument, ParseError, RawBlock
from .structure import is_chapter_line

_HEADING_STYLE = re.compile(r"heading\s*(\d)", re.IGNORECASE)


def parse_docx(path: Path) -> ParsedDocument:
    try:
        document = docx.Document(str(path))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ParseError(f"not a readable DOCX file: {exc}") from exc

    blocks: list[RawBlock] = []
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            blocks.append(_paragraph_block(item))
        elif isinstance(item, Table):
            blocks.extend(_table_blocks(item))

    has_styled_headings = any(b.kind in ("title", "heading") for b in blocks)
    if not has_styled_headings:
        for b in blocks:
            if b.kind == "paragraph" and is_chapter_line(b.text.strip()):
                b.kind, b.level = "heading", 1

    has_title_block = any(b.kind == "title" for b in blocks)
    core_title = None if has_title_block else (document.core_properties.title or None)
    return ParsedDocument("docx", blocks, metadata_title=core_title)


def _paragraph_block(p: Paragraph) -> RawBlock:
    style = (p.style.name if p.style is not None else "") or ""
    lowered = style.lower()
    kind: BlockKind = "paragraph"
    level: int | None = None
    if lowered == "title":
        kind, level = "title", 1
    elif (m := _HEADING_STYLE.match(lowered)) is not None:
        kind, level = "heading", min(max(int(m.group(1)), 1), 6)
    elif lowered.startswith("list") or _has_numbering(p):
        kind = "list_item"
    elif lowered in ("quote", "intense quote"):
        kind = "quote"
    return RawBlock(kind, p.text, level=level)


def _has_numbering(p: Paragraph) -> bool:
    ppr = p._p.pPr
    return ppr is not None and ppr.numPr is not None


def _table_blocks(table: Table) -> list[RawBlock]:
    blocks: list[RawBlock] = []
    seen: set[int] = set()
    for row in table.rows:
        for cell in row.cells:
            if id(cell._tc) in seen:  # merged cells repeat the same element
                continue
            seen.add(id(cell._tc))
            blocks.append(RawBlock("paragraph", "\n".join(p.text for p in cell.paragraphs)))
    return blocks
