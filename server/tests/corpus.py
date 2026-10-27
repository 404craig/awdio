"""Public-domain passages (Lewis Carroll, 1865; Jane Austen, 1813) and generators for test fixture files."""

from __future__ import annotations

from pathlib import Path

import docx
import pymupdf
from ebooklib import epub

ALICE_TITLE = "Alice's Adventures in Wonderland"
ALICE: list[tuple[str, list[str]]] = [
    ("CHAPTER I. Down the Rabbit-Hole", [
        "Alice was beginning to get very tired of sitting by her sister on the bank, and of having nothing to do: "
        "once or twice she had peeped into the book her sister was reading, but it had no pictures or conversations "
        "in it, “and what is the use of a book,” thought Alice “without pictures or conversations?”",
        "So she was considering in her own mind (as well as she could, for the hot day made her feel very sleepy "
        "and stupid), whether the pleasure of making a daisy-chain would be worth the trouble of getting up and "
        "picking the daisies, when suddenly a White Rabbit with pink eyes ran close by her.",
        "There was nothing so very remarkable in that; nor did Alice think it so very much out of the way to hear "
        "the Rabbit say to itself, “Oh dear! Oh dear! I shall be late!” (when she thought it over afterwards, it "
        "occurred to her that she ought to have wondered at this, but at the time it all seemed quite natural); but "
        "when the Rabbit actually took a watch out of its waistcoat-pocket, and looked at it, and then hurried on, "
        "Alice started to her feet, for it flashed across her mind that she had never before seen a rabbit with "
        "either a waistcoat-pocket, or a watch to take out of it, and burning with curiosity, she ran across the "
        "field after it, and fortunately was just in time to see it pop down a large rabbit-hole under the hedge.",
        "In another moment down went Alice after it, never once considering how in the world she was to get out again.",
    ]),
    ("CHAPTER II. The Pool of Tears", [
        "“Curiouser and curiouser!” cried Alice (she was so much surprised, that for the moment she quite forgot how "
        "to speak good English); “now I’m opening out like the largest telescope that ever was! Good-bye, feet!” "
        "(for when she looked down at her feet, they seemed to be almost out of sight, they were getting so far off).",
    ]),
    ("CHAPTER III. A Caucus-Race and a Long Tale", [
        "They were indeed a queer-looking party that assembled on the bank, and all of them were dripping wet, "
        "cross, and uncomfortable.",
    ]),
]
ALICE_BULLETS = ["the birds with draggled feathers", "the animals with their fur clinging close to them"]

PRIDE_TITLE = "Pride and Prejudice"
PRIDE: list[tuple[str, list[str]]] = [
    ("Chapter 1", [
        "It is a truth universally acknowledged, that a single man in possession of a good fortune, must be in "
        "want of a wife.",
        "However little known the feelings or views of such a man may be on his first entering a neighbourhood, "
        "this truth is so well fixed in the minds of the surrounding families, that he is considered the rightful "
        "property of some one or other of their daughters.",
        "“My dear Mr. Bennet,” said his lady to him one day, “have you heard that Netherfield Park is let at last?”",
        "Mr. Bennet replied that he had not.",
        "“But it is,” returned she; “for Mrs. Long has just been here, and she told me all about it.”",
    ]),
    ("Chapter 2", [
        "Mr. Bennet was among the earliest of those who waited on Mr. Bingley. He had always intended to visit "
        "him, though to the last always assuring his wife that he should not go; and till the evening after the "
        "visit was paid she had no knowledge of it.",
    ]),
]
PRIDE_QUOTE = "A lady's imagination is very rapid; it jumps from admiration to love, from love to matrimony in a moment."
PRIDE_LIST = ["Jane", "Elizabeth", "Mary"]

HYPHENATED_WORD = "conversations"  # broken across a line in the PDF fixture to test repair
RUNNING_HEADER = "Alice's Adventures in Wonderland"


def pride_paragraphs() -> list[str]:
    return [p for _, paras in PRIDE for p in paras]


# --- PDF -------------------------------------------------------------------------------------

def ascii_quotes(text: str) -> str:
    """The PDF fixture uses the built-in Helvetica, which cannot draw curly quotes."""
    return text.translate(str.maketrans({"“": '"', "”": '"', "’": "'", "‘": "'"}))


class _PdfWriter:
    """Tiny flowing layout: wraps words into lines, adds a running header and page-number footer."""

    def __init__(self, doc: pymupdf.Document, width: float = 420, height: float = 540) -> None:
        self.doc, self.width, self.height = doc, width, height
        self.left, self.right, self.top, self.bottom = 50, width - 50, 80, height - 70
        self.page: pymupdf.Page | None = None
        self.y = 0.0

    def new_page(self) -> None:
        self.page = self.doc.new_page(width=self.width, height=self.height)
        number = self.doc.page_count
        self.page.insert_text((self.left, 40), RUNNING_HEADER, fontsize=8, fontname="helv")
        self.page.insert_text((self.width / 2 - 5, self.height - 30), str(number), fontsize=8, fontname="helv")
        self.y = self.top

    def text(self, text: str, size: float, font: str = "helv", *, gap_after: float = 9,
             hyphenate: str | None = None, prefix: str = "") -> None:
        if self.page is None:
            self.new_page()
        width = self.right - self.left - (12 if prefix else 0)
        lines: list[str] = []
        current = ""
        for word in text.split(" "):
            candidate = f"{current} {word}".strip()
            if pymupdf.get_text_length(candidate, fontname=font, fontsize=size) <= width:
                current = candidate
                continue
            if hyphenate and word.strip('?,."') == hyphenate and current:
                head, tail = word[: len(word) // 2], word[len(word) // 2:]
                lines.append(f"{current} {head}-")
                current = tail
                hyphenate = None
                continue
            lines.append(current)
            current = word
        lines.append(current)
        pitch = size * 1.25
        for i, line in enumerate(lines):
            if self.y + pitch > self.bottom:
                self.new_page()
            assert self.page is not None
            x = self.left
            if prefix and i == 0:
                self.page.insert_text((x, self.y + size), prefix, fontsize=size, fontname=font)
            self.page.insert_text((x + (12 if prefix else 0), self.y + size), line, fontsize=size, fontname=font)
            self.y += pitch
        self.y += gap_after


def make_pdf(path: Path) -> Path:
    doc = pymupdf.open()
    w = _PdfWriter(doc)
    for ci, (heading, paragraphs) in enumerate(ALICE):
        w.new_page()
        if ci == 0:
            w.text(ALICE_TITLE, 22, "hebo", gap_after=24)
        w.text(heading, 15, "hebo", gap_after=12)
        for pi, para in enumerate(paragraphs):
            w.text(ascii_quotes(para), 11, hyphenate=HYPHENATED_WORD if (ci, pi) == (0, 0) else None)
        if ci == len(ALICE) - 1:
            for item in ALICE_BULLETS:
                w.text(item, 11, prefix="-", gap_after=4)
    doc.set_page_labels([{"startpage": 0, "prefix": "", "style": "r", "firstpagenum": 1},
                         {"startpage": 1, "prefix": "", "style": "D", "firstpagenum": 1}])
    doc.save(path)
    doc.close()
    return path


def make_blank_pdf(path: Path) -> Path:
    doc = pymupdf.open()
    page = doc.new_page()
    page.draw_rect(pymupdf.Rect(50, 50, 200, 200), color=(0, 0, 0))
    doc.save(path)
    doc.close()
    return path


# --- DOCX / EPUB / TXT / MD ------------------------------------------------------------------

def make_docx(path: Path) -> Path:
    d = docx.Document()
    d.add_heading(PRIDE_TITLE, level=0)
    for heading, paragraphs in PRIDE:
        d.add_heading(heading, level=1)
        for para in paragraphs:
            d.add_paragraph(para)
    d.add_paragraph(PRIDE_QUOTE, style="Quote")
    for name in PRIDE_LIST:
        d.add_paragraph(name, style="List Bullet")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Longbourn"
    table.rows[0].cells[1].text = "Netherfield"
    d.save(str(path))
    return path


def make_epub(path: Path) -> Path:
    book = epub.EpubBook()
    book.set_identifier("awdio-test-pride")
    book.set_title(PRIDE_TITLE)
    book.set_language("en")
    chapters = []
    for i, (heading, paragraphs) in enumerate(PRIDE, start=1):
        body = "".join(f"<p>{p}</p>" for p in paragraphs)
        if i == 1:
            body += f"<blockquote><p>{PRIDE_QUOTE}</p></blockquote>"
            body += "<ul>" + "".join(f"<li>{n}</li>" for n in PRIDE_LIST) + "</ul>"
            body += "<p>Kitty<br/>and Lydia</p><script>var hidden = 1;</script>"
        chapter = epub.EpubHtml(title=heading, file_name=f"chap_{i}.xhtml", lang="en")
        chapter.content = f"<html><head><title>{heading}</title></head><body><h1>{heading}</h1>{body}</body></html>"
        book.add_item(chapter)
        chapters.append(chapter)
    book.toc = chapters
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", *chapters]
    epub.write_epub(str(path), book)
    return path


def make_txt(path: Path) -> Path:
    parts = [PRIDE_TITLE.upper(), "By Jane Austen"]
    for heading, paragraphs in PRIDE:
        parts.append(heading)
        parts.extend(_wrap(p, 60) for p in paragraphs)
    path.write_text("\n\n".join(parts) + "\n", encoding="utf-8")
    return path


def make_md(path: Path) -> Path:
    lines = [f"# {PRIDE_TITLE}", "", "*By Jane Austen*", ""]
    for heading, paragraphs in PRIDE:
        lines += [f"## {heading}", ""]
        for p in paragraphs:
            lines += [_wrap(p, 70), ""]
    lines += [f"> {PRIDE_QUOTE}", ""]
    lines += [f"- **{n}** Bennet" for n in PRIDE_LIST] + [""]
    lines += ["See [the novel](https://www.gutenberg.org/ebooks/1342) and `chapter_one`.", ""]
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def _wrap(text: str, width: int) -> str:
    lines, current = [], ""
    for word in text.split(" "):
        if current and len(current) + 1 + len(word) > width:
            lines.append(current)
            current = word
        else:
            current = f"{current} {word}".strip()
    lines.append(current)
    return "\n".join(lines)
